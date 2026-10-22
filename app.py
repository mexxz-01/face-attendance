"""
=============================================================================
 Imperial Face Attendance System — High-Performance Python & FastAPI Backend
=============================================================================
Features:
  - Vectorized NumPy Euclidean matching (< 0.05ms)
  - SQLite persistent storage with automatic migration from legacy binary
  - Real-time Admin facial identification with instantaneous portal bypass
  - Server-side styled Excel (.xlsx) & PDF report generation
  - Secure Token authentication (Default Password: mexxz@2007)
=============================================================================
"""

import os
import json
import time
import secrets
import sqlite3
import struct
from datetime import datetime
from typing import List, Optional
from io import BytesIO

import numpy as np
from fastapi import FastAPI, Header, HTTPException, Request, Response
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# ── Configuration ──────────────────────────────────────────────────────────

PORT = 8080
ADMIN_PASSWORD = "mexxz@2007"
MATCH_THRESHOLD = 0.60
DESC_SIZE = 128
DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
WWW_DIR = os.path.join(os.path.dirname(__file__), "www")
DB_PATH = os.path.join(DATA_DIR, "attendance.db")
LEGACY_USERS = os.path.join(DATA_DIR, "users.dat")
LEGACY_ATTEND = os.path.join(DATA_DIR, "attend.dat")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(WWW_DIR, exist_ok=True)

# ── Database Initialization & Migration ─────────────────────────────────────

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            descriptor TEXT NOT NULL,
            is_admin INTEGER DEFAULT 0,
            created_at TEXT
        );
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            date TEXT NOT NULL,
            timestamp TEXT NOT NULL
        );
    """)
    conn.commit()

    # Migrate from legacy binary files if users table is empty
    cur.execute("SELECT COUNT(*) FROM users")
    count = cur.fetchone()[0]
    if count == 0 and os.path.exists(LEGACY_USERS):
        try:
            with open(LEGACY_USERS, "rb") as f:
                header = f.read(4)
                if len(header) == 4:
                    num_users = struct.unpack("i", header)[0]
                    fsize = os.path.getsize(LEGACY_USERS)
                    rec_size = (fsize - 4) // max(1, num_users)
                    for _ in range(num_users):
                        name_bytes = f.read(128)
                        name = name_bytes.split(b"\x00")[0].decode("utf-8", errors="ignore").strip()
                        desc_floats = struct.unpack(f"{DESC_SIZE}f", f.read(4 * DESC_SIZE))
                        is_adm = 0
                        if rec_size >= 644:
                            is_adm_bytes = f.read(4)
                            if len(is_adm_bytes) == 4:
                                is_adm = struct.unpack("i", is_adm_bytes)[0]
                        if name:
                            cur.execute("""
                                INSERT OR IGNORE INTO users (name, descriptor, is_admin, created_at)
                                VALUES (?, ?, ?, ?)
                            """, (name, json.dumps(list(desc_floats)), is_adm, datetime.now().isoformat()))
            conn.commit()
            print(f"[MIGRATION] Imported legacy users from {LEGACY_USERS}")
        except Exception as e:
            print(f"[MIGRATION_ERR] Users: {e}")

    # Migrate attendance
    cur.execute("SELECT COUNT(*) FROM attendance")
    att_count = cur.fetchone()[0]
    if att_count == 0 and os.path.exists(LEGACY_ATTEND):
        try:
            with open(LEGACY_ATTEND, "rb") as f:
                header = f.read(4)
                if len(header) == 4:
                    num_att = struct.unpack("i", header)[0]
                    for _ in range(num_att):
                        name_bytes = f.read(128)
                        name = name_bytes.split(b"\x00")[0].decode("utf-8", errors="ignore").strip()
                        date_bytes = f.read(16)
                        date_str = date_bytes.split(b"\x00")[0].decode("utf-8", errors="ignore").strip()
                        if name and date_str:
                            cur.execute("""
                                INSERT INTO attendance (name, date, timestamp)
                                VALUES (?, ?, ?)
                            """, (name, date_str, datetime.now().strftime("%H:%M:%S")))
            conn.commit()
            print(f"[MIGRATION] Imported legacy attendance from {LEGACY_ATTEND}")
        except Exception as e:
            print(f"[MIGRATION_ERR] Attendance: {e}")

    conn.close()

init_db()

# ── In-Memory Vector Cache for Sub-Millisecond Matching ─────────────────────

class BiometricRegistry:
    def __init__(self):
        self.names = []
        self.descriptors = None  # np.ndarray shape (N, 128)
        self.admins = set()
        self.reload()

    def reload(self):
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT name, descriptor, is_admin FROM users")
        rows = cur.fetchall()
        conn.close()

        names, descs, admins = [], [], set()
        for r in rows:
            names.append(r["name"])
            descs.append(json.loads(r["descriptor"]))
            if r["is_admin"] == 1:
                admins.add(r["name"])

        self.names = names
        self.descriptors = np.array(descs, dtype=np.float32) if descs else np.empty((0, DESC_SIZE), dtype=np.float32)
        self.admins = admins
        print(f"[REGISTRY] Loaded {len(self.names)} face embeddings into NumPy cache ({len(self.admins)} Admins)")

    def match(self, query_descriptor: List[float]):
        if len(self.names) == 0:
            return None, 999.0, False

        q = np.array(query_descriptor, dtype=np.float32)
        # Vectorized Euclidean distance calculation in C/BLAS via NumPy
        dists = np.linalg.norm(self.descriptors - q, axis=1)
        best_idx = int(np.argmin(dists))
        best_dist = float(dists[best_idx])

        if best_dist <= MATCH_THRESHOLD:
            name = self.names[best_idx]
            is_adm = name in self.admins
            return name, best_dist, is_adm

        return None, best_dist, False

registry = BiometricRegistry()

# ── Security Session Manager ────────────────────────────────────────────────

active_tokens = set()

def create_admin_token() -> str:
    token = secrets.token_hex(32)
    active_tokens.add(token)
    return token

def verify_token(token: Optional[str]) -> bool:
    if not token:
        return False
    if token.startswith("Bearer "):
        token = token[7:].strip()
    return token in active_tokens

# ── FastAPI App Setup ───────────────────────────────────────────────────────

app = FastAPI(title="Imperial Face Attendance Engine")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Pydantic Request Models ─────────────────────────────────────────────────

class RegisterRequest(BaseModel):
    name: str
    descriptor: List[float]

class VerifyRequest(BaseModel):
    descriptor: List[float]

class LoginRequest(BaseModel):
    password: str

class ToggleAdminRequest(BaseModel):
    name: str
    is_admin: int

class DeleteRequest(BaseModel):
    name: str

# ── API Routes: Public Biometric Recognition ────────────────────────────────

@app.post("/api/register")
def register_face(req: RegisterRequest):
    name = req.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Missing student name")
    if len(req.descriptor) != DESC_SIZE:
        raise HTTPException(status_code=400, detail=f"Expected {DESC_SIZE} descriptor values")

    today = datetime.now().strftime("%Y-%m-%d")
    now_time = datetime.now().strftime("%H:%M:%S")

    conn = get_db()
    cur = conn.cursor()

    # Insert or update user
    cur.execute("""
        INSERT INTO users (name, descriptor, is_admin, created_at)
        VALUES (?, ?, 0, ?)
        ON CONFLICT(name) DO UPDATE SET descriptor=excluded.descriptor
    """, (name, json.dumps(req.descriptor), datetime.now().isoformat()))

    # Check if already marked attendance today
    cur.execute("SELECT id FROM attendance WHERE name=? AND date=?", (name, today))
    if not cur.fetchone():
        cur.execute("INSERT INTO attendance (name, date, timestamp) VALUES (?, ?, ?)", (name, today, now_time))

    conn.commit()
    conn.close()

    registry.reload()
    print(f"[REG] Enrolled & marked present: {name}")

    return {
        "success": True,
        "message": "User enrolled",
        "name": name,
        "date": today,
        "is_admin": False
    }

@app.post("/api/verify")
def verify_face(req: VerifyRequest):
    if len(req.descriptor) != DESC_SIZE:
        raise HTTPException(status_code=400, detail="Invalid face descriptor length")

    name, dist, is_adm = registry.match(req.descriptor)
    today = datetime.now().strftime("%Y-%m-%d")

    if not name:
        return {
            "matched": False,
            "distance": round(dist, 4),
            "message": "Face not recognized"
        }

    # Record attendance
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM attendance WHERE name=? AND date=?", (name, today))
    already_marked = cur.fetchone() is not None

    if not already_marked:
        cur.execute("INSERT INTO attendance (name, date, timestamp) VALUES (?, ?, ?)",
                    (name, today, datetime.now().strftime("%H:%M:%S")))
        conn.commit()
    conn.close()

    admin_token = ""
    if is_adm:
        admin_token = create_admin_token()
        print(f"[ADMIN_BYPASS] Recognized Admin: {name} -> Issued session token")

    return {
        "matched": True,
        "name": name,
        "distance": round(dist, 4),
        "already_marked": already_marked,
        "date": today,
        "is_admin": is_adm,
        "admin_token": admin_token
    }

# ── API Routes: Secured Administrator Operations ────────────────────────────

@app.post("/api/admin/login")
def admin_login(req: LoginRequest):
    if req.password != ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Access Denied: Invalid Security Key")

    token = create_admin_token()
    print("[ADMIN] Password authenticated successfully")
    return {"success": True, "token": token}

@app.get("/api/admin/verify")
def admin_verify(authorization: Optional[str] = Header(None)):
    if not verify_token(authorization):
        raise HTTPException(status_code=401, detail="Unauthorized")
    return {"valid": True}

@app.get("/api/admin/attendance")
def admin_get_attendance(authorization: Optional[str] = Header(None)):
    if not verify_token(authorization):
        raise HTTPException(status_code=401, detail="Unauthorized")

    today = datetime.now().strftime("%Y-%m-%d")
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT name, is_admin FROM users ORDER BY name COLLATE NOCASE")
    all_users = cur.fetchall()

    cur.execute("SELECT name, timestamp FROM attendance WHERE date=?", (today,))
    present_map = {row["name"]: row["timestamp"] for row in cur.fetchall()}
    conn.close()

    user_list = []
    present_count = 0
    for u in all_users:
        name = u["name"]
        is_pres = name in present_map
        if is_pres:
            present_count += 1
        user_list.append({
            "name": name,
            "status": "present" if is_pres else "absent",
            "is_admin": bool(u["is_admin"]),
            "timestamp": present_map.get(name, "")
        })

    return {
        "date": today,
        "total": len(all_users),
        "present": present_count,
        "absent": len(all_users) - present_count,
        "users": user_list
    }

@app.post("/api/admin/toggle_admin")
def admin_toggle_role(req: ToggleAdminRequest, authorization: Optional[str] = Header(None)):
    if not verify_token(authorization):
        raise HTTPException(status_code=401, detail="Unauthorized")

    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE users SET is_admin=? WHERE name=?", (1 if req.is_admin else 0, req.name))
    updated = cur.rowcount > 0
    conn.commit()
    conn.close()

    if not updated:
        raise HTTPException(status_code=404, detail="Student not found")

    registry.reload()
    print(f"[ADMIN_ROLE] {req.name} set to {'Admin' if req.is_admin else 'Student'}")
    return {"success": True}

@app.post("/api/admin/delete")
def admin_delete_user(req: DeleteRequest, authorization: Optional[str] = Header(None)):
    if not verify_token(authorization):
        raise HTTPException(status_code=401, detail="Unauthorized")

    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM users WHERE name=?", (req.name,))
    cur.execute("DELETE FROM attendance WHERE name=?", (req.name,))
    conn.commit()
    conn.close()

    registry.reload()
    print(f"[ADMIN_DEL] Removed user: {req.name}")
    return {"success": True}

# ── API Routes: Server-Side Excel & PDF Reporting ───────────────────────────

@app.get("/api/admin/export/excel")
def export_excel(authorization: Optional[str] = Header(None)):
    if not verify_token(authorization):
        raise HTTPException(status_code=401, detail="Unauthorized")

    today = datetime.now().strftime("%Y-%m-%d")
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT name, is_admin FROM users ORDER BY name")
    users = cur.fetchall()
    cur.execute("SELECT name, timestamp FROM attendance WHERE date=?", (today,))
    present_map = {row["name"]: row["timestamp"] for row in cur.fetchall()}
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Styling Palette
    header_fill = PatternFill(start_color="4E0710", end_color="4E0710", fill_type="solid")
    gold_fill = PatternFill(start_color="D4AF37", end_color="D4AF37", fill_type="solid")
    alt_fill = PatternFill(start_color="F9F5F0", end_color="F9F5F0", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)

    # Title Banner
    ws.merge_cells("A1:E1")
    ws["A1"] = f"BIOMETRIC ATTENDANCE REPORT — {today}"
    ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="D4AF37")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    headers = ["#", "Student Name", "Privilege Role", "Status", "Recorded Time"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 24

    for idx, u in enumerate(users, 1):
        row_num = idx + 2
        name = u["name"]
        role = "Admin" if u["is_admin"] else "Student"
        is_pres = name in present_map
        status = "Present" if is_pres else "Absent"
        time_str = present_map.get(name, "—")

        ws.cell(row=row_num, column=1, value=idx).font = bold_font
        ws.cell(row=row_num, column=2, value=name).font = regular_font
        ws.cell(row=row_num, column=3, value=role).font = bold_font if u["is_admin"] else regular_font
        s_cell = ws.cell(row=row_num, column=4, value=status)
        s_cell.font = Font(name="Segoe UI", size=10, bold=True, color="008000" if is_pres else "CC0000")
        ws.cell(row=row_num, column=5, value=time_str).font = regular_font

        if idx % 2 == 0:
            for c in range(1, 6):
                ws.cell(row=row_num, column=c).fill = alt_fill

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 18

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Attendance_{today}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@app.get("/api/admin/export/pdf")
def export_pdf(authorization: Optional[str] = Header(None)):
    if not verify_token(authorization):
        raise HTTPException(status_code=401, detail="Unauthorized")

    today = datetime.now().strftime("%Y-%m-%d")
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT name, is_admin FROM users ORDER BY name")
    users = cur.fetchall()
    cur.execute("SELECT name, timestamp FROM attendance WHERE date=?", (today,))
    present_map = {row["name"]: row["timestamp"] for row in cur.fetchall()}
    conn.close()

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ImperialTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        textColor=colors.HexColor('#4E0710'),
        spaceAfter=6
    )
    sub_style = ParagraphStyle(
        'ImperialSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#555555'),
        spaceAfter=15
    )

    elements.append(Paragraph("BIOMETRIC ATTENDANCE REPORT", title_style))
    elements.append(Paragraph(f"Official Biometric Roster &bull; Date: <b>{today}</b> &bull; Total Enrolled: <b>{len(users)}</b> &bull; Present: <b>{len(present_map)}</b>", sub_style))

    data = [["#", "Student Name", "Privilege Role", "Status", "Time"]]
    for idx, u in enumerate(users, 1):
        name = u["name"]
        role = "ADMIN" if u["is_admin"] else "STUDENT"
        is_pres = name in present_map
        status = "PRESENT" if is_pres else "ABSENT"
        time_str = present_map.get(name, "—")
        data.append([str(idx), name, role, status, time_str])

    t = Table(data, colWidths=[30, 200, 100, 100, 100])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4E0710')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#F5E8BE')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E0D0C5')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FBF8F5')]),
    ]))
    elements.append(t)

    doc.build(elements)
    buf.seek(0)

    filename = f"Attendance_{today}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# ── Static Files & Single Page Routes ───────────────────────────────────────

@app.get("/admin")
def serve_admin():
    return FileResponse(os.path.join(WWW_DIR, "admin.html"))

@app.get("/")
def serve_kiosk():
    return FileResponse(os.path.join(WWW_DIR, "index.html"))

app.mount("/", StaticFiles(directory=WWW_DIR, html=True), name="static")

# ── Main Entrypoint ─────────────────────────────────────────────────────────

if __name__ == "__main__":
    import uvicorn
    print("=" * 60)
    print("  Imperial Face Attendance — Python FastAPI Engine")
    print(f"  Live Terminal: http://localhost:{PORT}")
    print(f"  Admin Portal:  http://localhost:{PORT}/admin")
    print(f"  Access Key:    {ADMIN_PASSWORD}")
    print("=" * 60)
    uvicorn.run("app:app", host="0.0.0.0", port=PORT, reload=False, log_level="info")
